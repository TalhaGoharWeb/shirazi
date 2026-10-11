"""
SHIRAZI plugin: presentation_designer.

Generates a .pptx presentation from a title + bullet outline (python-pptx,
free). Saves to ~/Documents/SHIRAZI/ and returns the path. Clean,
dark-futurist default styling that matches the SHIRAZI aesthetic.
"""

from __future__ import annotations

import os as _os
from pathlib import Path as _Path

PLUGIN = {
    "name": "presentation_designer",
    "description": (
        "Create a PowerPoint presentation from a title and slide outline. "
        "Trigger on 'make a presentation', 'create slides about <topic>', "
        "'presentation on <topic>'. Pass 'title' and 'slides' (list of "
        "{heading, bullets[]})."
    ),
    "parameters": {
        "type": "OBJECT",
        "properties": {
            "title": {"type": "STRING", "description": "Presentation title."},
            "slides": {
                "type": "ARRAY",
                "description": "Slides as {heading, bullets[]} objects.",
                "items": {"type": "OBJECT"},
            },
        },
        "required": ["title", "slides"],
    },
}

_BG = (6, 10, 16)
_TITLE_C = (0, 212, 255)
_BODY_C = (216, 240, 248)
_ACC_C = (255, 179, 71)


def run(parameters: dict, player=None, session_memory=None) -> str:
    title = str(parameters.get("title") or "Untitled").strip()[:120]
    slides = parameters.get("slides") or []
    if not isinstance(slides, list) or not slides:
        return "Sir, I need at least one slide with a heading and bullets."

    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
        from pptx.dml.color import RGBColor
        from pptx.enum.text import PP_ALIGN
    except Exception:
        return ("Sir, python-pptx is not installed. "
                "Run: pip install python-pptx")

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    def paint_bg(slide):
        bg = slide.background
        fill = bg.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor(*_BG)

    def add_textbox(slide, left, top, width, height):
        return slide.shapes.add_textbox(
            Inches(left), Inches(top), Inches(width), Inches(height))

    # Title slide
    s = prs.slides.add_slide(blank)
    paint_bg(s)
    tb = add_textbox(s, 1, 2.4, 11.3, 2.0)
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = RGBColor(*_TITLE_C)
    p.alignment = PP_ALIGN.CENTER

    # Content slides
    for sl in slides[:20]:
        if not isinstance(sl, dict):
            continue
        heading = str(sl.get("heading") or "")[:120]
        bullets = sl.get("bullets") or []
        if not heading:
            continue
        s = prs.slides.add_slide(blank)
        paint_bg(s)
        tb = add_textbox(s, 0.8, 0.4, 11.7, 1.2)
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = heading
        p.font.size = Pt(32)
        p.font.bold = True
        p.font.color.rgb = RGBColor(*_ACC_C)
        tb = add_textbox(s, 1.0, 1.9, 11.3, 5.0)
        tf = tb.text_frame
        tf.word_wrap = True
        first = True
        for b in bullets[:8]:
            b = str(b)[:300]
            if not b:
                continue
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.text = b
            p.level = 0
            p.font.size = Pt(20)
            p.font.color.rgb = RGBColor(*_BODY_C)
            p.space_after = Pt(10)

    out_dir = _Path(_os.path.expanduser("~")) / "Documents" / "SHIRAZI"
    out_dir.mkdir(parents=True, exist_ok=True)
    safe = "".join(c if c.isalnum() or c in " -_" else "_" for c in title)[:60]
    out = out_dir / f"{safe or 'presentation'}.pptx"
    i = 1
    while out.exists():
        out = out_dir / f"{safe or 'presentation'} ({i}).pptx"
        i += 1
    try:
        prs.save(str(out))
    except Exception as e:
        return f"Sir, saving the presentation failed: {e}"
    text = f"Presentation ready, sir: {out}"
    if player:
        try:
            player.write_log(f"SHIRAZI: {text}")
        except Exception:
            pass
    return text
